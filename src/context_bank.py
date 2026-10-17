from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping


SCHEMA_VERSION = "2026-10-17.1"
HEADER_ROW = 2
DATA_START_ROW = 3
CONTEXT_DELIMITER = "||"
TAG_DELIMITER = ";"
ID_NAMESPACE = uuid.UUID("2ca25cfe-5ca8-4d2a-b5e9-2ddce5664342")

EXPECTED_HEADERS = (
    "Vocab Source",
    "English Vocab",
    "English Definition",
    "English Context (MICEBooks)",
    "ASL Database Video",
    "ASL EXACT Syn/Homonyms",
    "ASL NEAR Syn/Homonyms",
    "Recording Status",
    "DPAN Notes",
)

STATUS_CODES = {
    "": "missing",
    "-": "default",
    "remove this": "remove",
    "need to film": "needs_filming",
    "use this": "use",
    "duplicate word(s)": "duplicate",
    "in 562 dataset": "existing_reference",
}
ACTIVE_STATUS_CODES = {"default", "use", "existing_reference"}


class ContextBankError(ValueError):
    """Raised when a workbook cannot be safely transformed."""


@dataclass(frozen=True)
class WorkbookRow:
    source_row_number: int
    source: str
    word: str
    definition: str
    context: str
    video_ref: str
    exact_tags: str
    near_tags: str
    recording_status: str
    notes: str

    @property
    def values(self) -> tuple[str, ...]:
        return (
            self.source,
            self.word,
            self.definition,
            self.context,
            self.video_ref,
            self.exact_tags,
            self.near_tags,
            self.recording_status,
            self.notes,
        )


@dataclass(frozen=True)
class ContextBank:
    source_name: str
    workbook_sha256: str
    sheet_name: str
    rows: tuple[WorkbookRow, ...]


@dataclass(frozen=True)
class ImportPayload:
    import_run_id: str
    workbook_sha256: str
    tables: dict[str, list[dict[str, Any]]]
    summary: dict[str, Any]


def raw_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", raw_text(value)).strip()


def key_text(value: object) -> str:
    return clean_text(value).casefold()


def split_values(value: object, delimiter: str) -> list[tuple[str, str]]:
    raw_value = raw_text(value)
    if not raw_value.strip():
        return []

    result: list[tuple[str, str]] = []
    for part in raw_value.split(delimiter):
        raw_part = part.strip()
        normalized = clean_text(raw_part)
        if normalized:
            result.append((raw_part, normalized))
    return result


def recording_status_code(value: object) -> str:
    status_key = key_text(value)
    try:
        return STATUS_CODES[status_key]
    except KeyError as exc:
        raise ContextBankError(f"Unknown recording status: {clean_text(value)!r}") from exc


def video_ref_format(value: object) -> str:
    text = clean_text(value)
    if re.fullmatch(r".+\.\d+\.\d+", text):
        return "base_number_number"
    if re.fullmatch(r".+\.\d+", text):
        return "base_number"
    if text and "." not in text:
        return "bare_legacy_key"
    return "unrecognized"


def stable_uuid(kind: str, *parts: object) -> str:
    normalized = "\x1f".join(str(part) for part in parts)
    return str(uuid.uuid5(ID_NAMESPACE, f"{kind}\x1e{normalized}"))


def canonical_video_key(variant_id: str) -> str:
    return f"psv_{uuid.UUID(variant_id).hex}"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rows_from_values(rows: Iterable[tuple[int, Iterable[object]]]) -> tuple[WorkbookRow, ...]:
    parsed: list[WorkbookRow] = []
    for source_row_number, values in rows:
        cells = tuple(raw_text(value) for value in values)
        if len(cells) != len(EXPECTED_HEADERS):
            raise ContextBankError(
                f"Row {source_row_number} has {len(cells)} columns; "
                f"expected {len(EXPECTED_HEADERS)}."
            )
        if not any(cell.strip() for cell in cells):
            continue
        parsed.append(WorkbookRow(source_row_number, *cells))
    return tuple(parsed)


def load_context_bank(path: Path, sheet_name: str | None = None) -> ContextBank:
    from openpyxl import load_workbook

    resolved_path = path.expanduser().resolve()
    if not resolved_path.is_file():
        raise FileNotFoundError(f"Context bank workbook not found: {resolved_path}")

    workbook = load_workbook(resolved_path, read_only=True, data_only=False)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        header_values = next(
            worksheet.iter_rows(
                min_row=HEADER_ROW,
                max_row=HEADER_ROW,
                min_col=1,
                max_col=len(EXPECTED_HEADERS),
                values_only=True,
            )
        )
        headers = tuple(clean_text(value) for value in header_values)
        if headers != EXPECTED_HEADERS:
            raise ContextBankError(
                "The workbook columns do not match the audited PopSign context bank. "
                f"Expected {EXPECTED_HEADERS!r}; received {headers!r}."
            )

        values = (
            (row_number, row)
            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=DATA_START_ROW,
                    min_col=1,
                    max_col=len(EXPECTED_HEADERS),
                    values_only=True,
                ),
                start=DATA_START_ROW,
            )
        )
        parsed_rows = rows_from_values(values)
        return ContextBank(
            source_name=resolved_path.name,
            workbook_sha256=file_sha256(resolved_path),
            sheet_name=worksheet.title,
            rows=parsed_rows,
        )
    finally:
        workbook.close()


def _row_fingerprint(row: WorkbookRow) -> str:
    serialized = json.dumps(
        [row.source_row_number, *row.values],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def _variant_review_status(format_name: str, has_definition: bool) -> str:
    if format_name == "unrecognized":
        return "needs_video_key_review"
    if not has_definition:
        return "needs_definition"
    return "unreviewed"


def build_import_payload(bank: ContextBank) -> ImportPayload:
    if not bank.rows:
        raise ContextBankError("The context bank contains no populated data rows.")

    import_run_id = stable_uuid(
        "import_run", SCHEMA_VERSION, bank.workbook_sha256, bank.sheet_name
    )

    prepared: list[dict[str, Any]] = []
    pair_first_source_row: dict[tuple[str, str], str] = {}
    word_display: dict[str, str] = {}
    variant_aggregate: dict[tuple[str, str], dict[str, Any]] = {}
    phrase_display: dict[str, str] = {}
    phrase_assignment_count: Counter[str] = Counter()
    word_phrase_assignment_count: Counter[tuple[str, str]] = Counter()
    active_phrase_variants: dict[tuple[str, str], set[str]] = defaultdict(set)
    all_phrase_variants: dict[tuple[str, str], set[str]] = defaultdict(set)

    for row in bank.rows:
        source = clean_text(row.source)
        word = clean_text(row.word)
        word_key = word.casefold()
        video_ref = clean_text(row.video_ref)
        video_key_normalized = video_ref.casefold()
        if not source:
            raise ContextBankError(f"Source is missing on workbook row {row.source_row_number}.")
        if not word:
            raise ContextBankError(f"Word is missing on workbook row {row.source_row_number}.")
        if not video_ref:
            raise ContextBankError(
                f"Video reference is missing on workbook row {row.source_row_number}."
            )

        status_code = recording_status_code(row.recording_status)
        active = status_code in ACTIVE_STATUS_CODES
        word_id = stable_uuid("word", word_key)
        variant_id = stable_uuid("variant", word_key, video_key_normalized)
        source_row_id = stable_uuid("source_row", import_run_id, row.source_row_number)
        pair_key = (word_key, video_key_normalized)
        duplicate_of = pair_first_source_row.get(pair_key)
        pair_first_source_row.setdefault(pair_key, source_row_id)

        contexts = split_values(row.context, CONTEXT_DELIMITER)
        exact_tags = split_values(row.exact_tags, TAG_DELIMITER)
        near_tags = split_values(row.near_tags, TAG_DELIMITER)
        format_name = video_ref_format(video_ref)

        word_display.setdefault(word_key, word)
        aggregate = variant_aggregate.setdefault(
            pair_key,
            {
                "id": variant_id,
                "word_id": word_id,
                "video_key": canonical_video_key(variant_id),
                "legacy_video_ref": video_ref,
                "video_ref_key": video_key_normalized,
                "video_ref_format": format_name,
                "status_codes": set(),
                "source_row_count": 0,
                "is_active": False,
                "has_definition": False,
            },
        )
        aggregate["status_codes"].add(status_code)
        aggregate["source_row_count"] += 1
        aggregate["is_active"] = bool(aggregate["is_active"] or active)
        aggregate["has_definition"] = bool(
            aggregate["has_definition"] or clean_text(row.definition)
        )

        prepared_row = {
            "source": row,
            "source_id": source_row_id,
            "word": word,
            "word_key": word_key,
            "word_id": word_id,
            "variant_id": variant_id,
            "video_ref": video_ref,
            "video_ref_key": video_key_normalized,
            "video_ref_format": format_name,
            "status_code": status_code,
            "active": active,
            "duplicate_of": duplicate_of,
            "contexts": contexts,
            "exact_tags": exact_tags,
            "near_tags": near_tags,
        }
        prepared.append(prepared_row)

        for _, context_text in contexts:
            phrase_key = context_text.casefold()
            phrase_display.setdefault(phrase_key, context_text)
            phrase_assignment_count[phrase_key] += 1
            word_phrase_assignment_count[(word_key, phrase_key)] += 1
            all_phrase_variants[(word_key, phrase_key)].add(variant_id)
            if active:
                active_phrase_variants[(word_key, phrase_key)].add(variant_id)

    words = [
        {
            "id": stable_uuid("word", word_key),
            "word_key": word_key,
            "display_word": display_word,
            "review_status": "unreviewed",
        }
        for word_key, display_word in word_display.items()
    ]

    variants: list[dict[str, Any]] = []
    for aggregate in variant_aggregate.values():
        statuses = sorted(aggregate.pop("status_codes"))
        has_definition = bool(aggregate.pop("has_definition"))
        format_name = str(aggregate["video_ref_format"])
        is_active = bool(aggregate["is_active"])
        variants.append(
            {
                **aggregate,
                "recording_status": statuses[0] if len(statuses) == 1 else "mixed",
                "source_statuses": statuses,
                "sense_label": None,
                "definition": None,
                "review_status": _variant_review_status(format_name, has_definition),
                "human_approved": False,
                "matching_eligible": is_active and format_name != "unrecognized",
            }
        )

    source_rows: list[dict[str, Any]] = []
    phrases = [
        {
            "id": stable_uuid("phrase", phrase_key),
            "phrase_key": phrase_key,
            "context_text": phrase_display[phrase_key],
            "source_assignment_count": phrase_assignment_count[phrase_key],
        }
        for phrase_key in phrase_display
    ]
    word_phrases = [
        {
            "id": stable_uuid("word_phrase", word_key, phrase_key),
            "word_id": stable_uuid("word", word_key),
            "phrase_id": stable_uuid("phrase", phrase_key),
            "source_assignment_count": count,
        }
        for (word_key, phrase_key), count in word_phrase_assignment_count.items()
    ]
    variant_contexts: list[dict[str, Any]] = []
    tags_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    variant_tags: list[dict[str, Any]] = []

    for item in prepared:
        row: WorkbookRow = item["source"]
        source_rows.append(
            {
                "id": item["source_id"],
                "import_run_id": import_run_id,
                "source_row_number": row.source_row_number,
                "row_fingerprint": _row_fingerprint(row),
                "word_id": item["word_id"],
                "variant_id": item["variant_id"],
                "raw_source": row.source,
                "source_key": key_text(row.source),
                "raw_word": row.word,
                "word_key": item["word_key"],
                "raw_definition": row.definition,
                "normalized_definition": clean_text(row.definition) or None,
                "raw_context": row.context,
                "raw_video_ref": row.video_ref,
                "video_ref_key": item["video_ref_key"],
                "raw_exact_tags": row.exact_tags,
                "raw_near_tags": row.near_tags,
                "raw_recording_status": row.recording_status,
                "recording_status": item["status_code"],
                "raw_notes": row.notes,
                "normalized_notes": clean_text(row.notes) or None,
                "is_active_candidate": item["active"],
                "context_count": len(item["contexts"]),
                "exact_tag_count": len(item["exact_tags"]),
                "near_tag_count": len(item["near_tags"]),
                "is_duplicate_pair": item["duplicate_of"] is not None,
                "duplicate_of_source_row_id": item["duplicate_of"],
            }
        )

        for position, (raw_context, context_text) in enumerate(item["contexts"], start=1):
            phrase_key = context_text.casefold()
            shared_all = len(all_phrase_variants[(item["word_key"], phrase_key)]) > 1
            shared_active = (
                len(active_phrase_variants[(item["word_key"], phrase_key)]) > 1
            )
            if not item["active"]:
                eligibility_reason = "inactive_source"
            elif item["video_ref_format"] == "unrecognized":
                eligibility_reason = "video_key_review"
            elif shared_active:
                eligibility_reason = "shared_across_active_variants"
            else:
                eligibility_reason = "unique_to_active_variant"
            matching_eligible = eligibility_reason == "unique_to_active_variant"
            variant_context_id = stable_uuid(
                "variant_context",
                item["source_id"],
                position,
                item["word_key"],
                phrase_key,
            )
            variant_contexts.append(
                {
                    "id": variant_context_id,
                    "variant_id": item["variant_id"],
                    "word_phrase_id": stable_uuid(
                        "word_phrase", item["word_key"], phrase_key
                    ),
                    "source_row_id": item["source_id"],
                    "context_position": position,
                    "raw_context_text": raw_context,
                    "shared_across_variants": shared_all,
                    "shared_across_active_variants": shared_active,
                    "matching_eligible": matching_eligible,
                    "eligibility_reason": eligibility_reason,
                    "review_status": (
                        "auto_screened" if matching_eligible else "needs_review"
                    ),
                    "human_approved": False,
                }
            )

        for tag_kind, tag_values in (
            ("exact", item["exact_tags"]),
            ("near", item["near_tags"]),
        ):
            for position, (raw_tag, display_tag) in enumerate(tag_values, start=1):
                tag_key = display_tag.casefold()
                compound_key = (tag_kind, tag_key)
                tag_id = stable_uuid("tag", tag_kind, tag_key)
                tags_by_key.setdefault(
                    compound_key,
                    {
                        "id": tag_id,
                        "tag_kind": tag_kind,
                        "tag_key": tag_key,
                        "display_tag": display_tag,
                        "review_status": "unreviewed",
                        "human_approved": False,
                    },
                )
                variant_tags.append(
                    {
                        "id": stable_uuid(
                            "variant_tag",
                            item["source_id"],
                            tag_kind,
                            position,
                            tag_key,
                        ),
                        "variant_id": item["variant_id"],
                        "tag_id": tag_id,
                        "source_row_id": item["source_id"],
                        "tag_position": position,
                        "raw_tag_text": raw_tag,
                        "review_status": "unreviewed",
                        "human_approved": False,
                    }
                )

    active_variant_ids = {row["id"] for row in variants if row["is_active"]}
    eligible_contexts_by_variant: dict[str, int] = Counter(
        row["variant_id"] for row in variant_contexts if row["matching_eligible"]
    )
    active_variants_by_word: dict[str, set[str]] = defaultdict(set)
    for variant in variants:
        if variant["is_active"]:
            active_variants_by_word[variant["word_id"]].add(variant["id"])
    ready_words = 0
    for active_variants in active_variants_by_word.values():
        if active_variants and all(
            eligible_contexts_by_variant.get(variant_id, 0) > 0
            for variant_id in active_variants
        ):
            ready_words += 1

    exact_assignment_count = sum(len(item["exact_tags"]) for item in prepared)
    near_assignment_count = sum(len(item["near_tags"]) for item in prepared)
    status_counts = Counter(
        clean_text(item["source"].recording_status) or "blank" for item in prepared
    )
    summary = {
        "schema_version": SCHEMA_VERSION,
        "workbook_sha256": bank.workbook_sha256,
        "source_rows": len(source_rows),
        "words": len(words),
        "sign_variants": len(variants),
        "phrases": len(phrases),
        "word_phrases": len(word_phrases),
        "variant_contexts": len(variant_contexts),
        "tags": len(tags_by_key),
        "variant_tags": len(variant_tags),
        "exact_tag_assignments": exact_assignment_count,
        "near_tag_assignments": near_assignment_count,
        "active_source_rows": sum(item["active"] for item in prepared),
        "active_sign_variants": len(active_variant_ids),
        "matching_eligible_contexts": sum(
            row["matching_eligible"] for row in variant_contexts
        ),
        "blocked_shared_active_contexts": sum(
            row["shared_across_active_variants"] for row in variant_contexts
        ),
        "words_matching_ready": ready_words,
        "recording_status_counts": dict(status_counts),
    }

    tables = {
        "popsign_import_runs": [
            {
                "id": import_run_id,
                "schema_version": SCHEMA_VERSION,
                "source_name": bank.source_name,
                "worksheet_name": bank.sheet_name,
                "workbook_sha256": bank.workbook_sha256,
                "source_row_count": len(source_rows),
                "status": "prepared",
                "phase": "base_data",
                "progress": {},
                "payload_counts": summary,
            }
        ],
        "popsign_words": words,
        "popsign_sign_variants": variants,
        "popsign_source_rows": source_rows,
        "popsign_phrases": phrases,
        "popsign_word_phrases": word_phrases,
        "popsign_variant_contexts": variant_contexts,
        "popsign_tags": list(tags_by_key.values()),
        "popsign_variant_tags": variant_tags,
    }
    payload = ImportPayload(
        import_run_id=import_run_id,
        workbook_sha256=bank.workbook_sha256,
        tables=tables,
        summary=summary,
    )
    validate_import_payload(payload)
    return payload


def validate_import_payload(payload: ImportPayload) -> None:
    tables = payload.tables
    id_tables = (
        "popsign_import_runs",
        "popsign_words",
        "popsign_sign_variants",
        "popsign_source_rows",
        "popsign_phrases",
        "popsign_word_phrases",
        "popsign_variant_contexts",
        "popsign_tags",
        "popsign_variant_tags",
    )
    for table in id_tables:
        ids = [row["id"] for row in tables[table]]
        if len(ids) != len(set(ids)):
            raise ContextBankError(f"{table} contains duplicate primary IDs.")

    word_ids = {row["id"] for row in tables["popsign_words"]}
    variant_ids = {row["id"] for row in tables["popsign_sign_variants"]}
    source_row_ids = {row["id"] for row in tables["popsign_source_rows"]}
    phrase_ids = {row["id"] for row in tables["popsign_phrases"]}
    word_phrase_ids = {row["id"] for row in tables["popsign_word_phrases"]}
    tag_ids = {row["id"] for row in tables["popsign_tags"]}

    if any(row["word_id"] not in word_ids for row in tables["popsign_sign_variants"]):
        raise ContextBankError("A sign variant references an unknown word.")
    if any(
        row["word_id"] not in word_ids or row["variant_id"] not in variant_ids
        for row in tables["popsign_source_rows"]
    ):
        raise ContextBankError("A source row has an invalid word or variant reference.")
    if any(
        row["word_id"] not in word_ids or row["phrase_id"] not in phrase_ids
        for row in tables["popsign_word_phrases"]
    ):
        raise ContextBankError("A word-phrase row has an invalid reference.")
    if any(
        row["variant_id"] not in variant_ids
        or row["word_phrase_id"] not in word_phrase_ids
        or row["source_row_id"] not in source_row_ids
        for row in tables["popsign_variant_contexts"]
    ):
        raise ContextBankError("A variant-context row has an invalid reference.")
    if any(
        row["variant_id"] not in variant_ids
        or row["tag_id"] not in tag_ids
        or row["source_row_id"] not in source_row_ids
        for row in tables["popsign_variant_tags"]
    ):
        raise ContextBankError("A variant-tag row has an invalid reference.")

    context_pairs = [
        (row["variant_id"], row["word_phrase_id"])
        for row in tables["popsign_variant_contexts"]
    ]
    if len(context_pairs) != len(set(context_pairs)):
        raise ContextBankError(
            "A variant has the same word-specific phrase assignment more than once."
        )

    variants = {row["id"]: row for row in tables["popsign_sign_variants"]}
    for row in tables["popsign_variant_contexts"]:
        if row["matching_eligible"]:
            if row["eligibility_reason"] != "unique_to_active_variant":
                raise ContextBankError("Eligible context has an invalid eligibility reason.")
            variant = variants[row["variant_id"]]
            if not variant["is_active"] or not variant["matching_eligible"]:
                raise ContextBankError("Eligible context belongs to an ineligible variant.")

    summary_table_keys = {
        "source_rows": "popsign_source_rows",
        "words": "popsign_words",
        "sign_variants": "popsign_sign_variants",
        "phrases": "popsign_phrases",
        "word_phrases": "popsign_word_phrases",
        "variant_contexts": "popsign_variant_contexts",
        "tags": "popsign_tags",
        "variant_tags": "popsign_variant_tags",
    }
    for summary_key, table in summary_table_keys.items():
        if payload.summary[summary_key] != len(tables[table]):
            raise ContextBankError(
                f"Summary count {summary_key} does not match {table}."
            )


def reconcile_with_audit(
    payload: ImportPayload,
    audit: Mapping[str, Any],
) -> None:
    expected = {
        "source_rows": audit["source"]["data_rows"],
        "words": audit["coverage"]["unique_words_case_insensitive"],
        "sign_variants": audit["coverage"]["unique_video_keys_case_insensitive"],
        "phrases": audit["coverage"]["unique_context_strings_case_insensitive"],
        "word_phrases": audit["coverage"]["unique_word_context_pairs"],
        "variant_contexts": audit["coverage"]["context_phrase_assignments"],
        "exact_tag_assignments": audit["coverage"]["exact_tag_assignments"],
        "near_tag_assignments": audit["coverage"]["near_tag_assignments"],
        "active_source_rows": audit["existing_active_status_policy"]["active_rows"],
        "active_sign_variants": audit["existing_active_status_policy"][
            "active_unique_video_keys"
        ],
    }
    mismatches = {
        key: {"expected": expected_value, "actual": payload.summary.get(key)}
        for key, expected_value in expected.items()
        if payload.summary.get(key) != expected_value
    }

    expected_hash = audit["source"]["sha256"]
    if payload.workbook_sha256 != expected_hash:
        mismatches["workbook_sha256"] = {
            "expected": expected_hash,
            "actual": payload.workbook_sha256,
        }

    if payload.summary["recording_status_counts"] != audit["recording_status_counts"]:
        mismatches["recording_status_counts"] = {
            "expected": audit["recording_status_counts"],
            "actual": payload.summary["recording_status_counts"],
        }

    if mismatches:
        raise ContextBankError(
            "The import payload does not reconcile with the Phase 1 audit: "
            + json.dumps(mismatches, sort_keys=True)
        )


def public_summary(payload: ImportPayload) -> dict[str, Any]:
    """Return counts and hashes only; no private phrases or identifiers."""

    return dict(payload.summary)
